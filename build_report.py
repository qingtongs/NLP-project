"""Fill the course DOCX template from a reviewed JSON report payload."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph


HEADINGS = (
    "Introduction", "Methodology", "Evaluation Result", "Analysis", "Contributions", "Conclusion"
)


def after(paragraph: Paragraph, text: str) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    new = Paragraph(element, paragraph._parent)
    new.style = "Normal"
    new.text = text
    new.alignment = (WD_ALIGN_PARAGRAPH.LEFT if text.startswith("References:")
                     else WD_ALIGN_PARAGRAPH.JUSTIFY)
    new.paragraph_format.space_after = Pt(6)
    for run in new.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0, 0, 0)
    return new


def set_body(paragraph: Paragraph, text: str) -> None:
    paragraph.text = text
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(6)
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
        run.font.italic = False
        run.font.color.rgb = RGBColor(0, 0, 0)


def remove(paragraph: Paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def grid_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), "000000")
        borders.append(tag)
    properties.append(borders)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--payload", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    payload = json.loads(args.payload.read_text(encoding="utf-8"))
    doc = Document(args.template)
    if "Title" not in [style.name for style in doc.styles]:
        title_style = doc.styles.add_style("Title", WD_STYLE_TYPE.PARAGRAPH)
        title_style.base_style = doc.styles["Normal"]
        title_style.font.name = "Arial"
        title_style.font.size = Pt(20)
        title_style.font.bold = True
        title_style.font.color.rgb = RGBColor(0, 0, 0)
    # Keep the template artwork, section geometry and remaining student table.
    title = next(p for p in doc.paragraphs if "Natural Language Processing" in p.text)
    title.text = "Natural Language Processing\nGroup 07 MultiHop RAG Project Report"
    title.style = "Title"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.name = "Arial"
        run.font.size = Pt(20)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
    date = next(p for p in doc.paragraphs if "2026" in p.text)
    date.text = payload["date"]
    for paragraph in list(doc.paragraphs):
        if "________" in paragraph.text:
            remove(paragraph)
    table = doc.tables[0]
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = payload["member_name"]
    table.cell(0, 2).text = "学号"
    table.cell(0, 3).text = payload["student_id"]
    for row in (table.rows[2], table.rows[1]):
        table._tbl.remove(row._tr)
    table.cell(1, 0).text = "专业班级"
    table.cell(1, 1).text = payload.get("class_name") or "待定"
    table.cell(1, 2).text = "课程代码"
    table.cell(1, 3).text = payload.get("course_code") or "待定"
    table.cell(2, 0).text = "课程名称"
    table.cell(2, 1).text = "Natural Language Processing"
    table.cell(2, 2).text = "任课教师"
    table.cell(2, 3).text = payload.get("teacher") or "待定"
    repository = next(p for p in doc.paragraphs if p.text.startswith("GitHub Repository:"))
    repository.text = "GitHub Repository: " + (payload.get("github_url") or "Pending publication")
    repository.paragraph_format.page_break_before = True
    for run in repository.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0, 0, 0)
    for heading in HEADINGS:
        target = next(p for p in doc.paragraphs if p.text.strip().startswith(heading))
        target.text = heading
        for run in target.runs:
            run.font.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0, 0, 0)
        placeholder = target._p.getnext()
        slot = Paragraph(placeholder, target._parent)
        texts = payload["sections"][heading]
        if not texts:
            raise ValueError(f"Missing section: {heading}")
        set_body(slot, texts[0])
        current = slot
        for paragraph_text in texts[1:]:
            current = after(current, paragraph_text)
        if heading == "Evaluation Result":
            scores = payload["scores"]
            score_table = doc.add_table(rows=1, cols=3)
            grid_borders(score_table)
            score_table.rows[0].cells[0].text = "Evaluation scope"
            score_table.rows[0].cells[1].text = "Questions"
            score_table.rows[0].cells[2].text = "Official Accuracy"
            for key in ("overall", "inference_query", "comparison_query", "temporal_query", "null_query"):
                row = score_table.add_row().cells
                row[0].text = key
                row[1].text = str(scores[key]["count"])
                row[2].text = f"{scores[key]['accuracy'] * 100:.2f}%"
            current._p.addnext(score_table._tbl)
            for row in score_table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.name = "Times New Roman"
                            run.font.size = Pt(10)
    for paragraph in list(doc.paragraphs):
        if paragraph.text.startswith("[Submission:") or paragraph.text.startswith("[Page Limit:") or paragraph.text.startswith("[Reminder:"):
            remove(paragraph)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    generated_path = args.output.with_suffix(".generated.docx")
    doc.save(generated_path)
    # python-docx reserializes unrelated package parts. Restore each untouched
    # part from the supplied reference so drawings and custom XML are exact.
    modified_parts = {"word/document.xml", "word/styles.xml"}
    with ZipFile(args.template) as source, ZipFile(generated_path) as generated, \
            ZipFile(args.output, "w") as final:
        if set(source.namelist()) != set(generated.namelist()):
            raise ValueError("Template package-part inventory changed")
        for info in source.infolist():
            data = (generated if info.filename in modified_parts else source).read(info.filename)
            final.writestr(info, data)
    generated_path.unlink()
    print(args.output)


if __name__ == "__main__":
    main()
